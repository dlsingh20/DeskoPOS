"""
Reading a client's menu sheet (.xlsx or .csv) into categories and items.

Restaurants send menus in many shapes, so the reader looks for the header row by name
(Name / Item, Category, Price / Rate / MRP, Veg ...). When there is no category column it
uses the sheet name, or a row that has a name but no price as a category heading
("STARTERS" followed by its dishes).
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

COLUMNS = {
    'name': ['name', 'item', 'item name', 'dish', 'dish name', 'product', 'product name', 'menu item', 'particulars'],
    'category': ['category', 'section', 'group', 'course', 'menu section', 'category name'],
    'price': ['price', 'rate', 'mrp', 'selling price', 'amount', 'price (rs)', 'price (inr)', 'rate (rs)'],
    'description': ['description', 'details', 'desc'],
    'vegetarian': ['vegetarian', 'veg', 'veg/non-veg', 'veg / non veg', 'veg/non veg', 'veg or non-veg', 'food type', 'type'],
    'vegan': ['vegan'],
    'gluten_free': ['gluten free', 'gluten-free'],
    'spicy': ['spicy'],
    'available': ['available', 'availability'],
    'allergens': ['allergens', 'allergen'],
    'calories': ['calories', 'kcal'],
    'prep_minutes': ['prep minutes', 'preparation minutes', 'prep time', 'preparation time'],
    'kitchen_notes': ['kitchen notes'],
    'short_name': ['short name', 'kot name'],
    'sort_order': ['display order', 'sort order'],
}
YES = {'yes', 'y', 'true', '1', 'veg', 'v', 'vegetarian', 'pure veg'}
NO = {'no', 'n', 'false', '0', 'non-veg', 'nonveg', 'non veg', 'nv', 'non-vegetarian', 'non vegetarian', 'egg'}
MAX_ITEMS = 3000
MAX_PROBLEMS = 20


class MenuSheetError(Exception):
    """The file can't be read as a menu; the message says why, for the person importing."""


def _key(text):
    return re.sub(r'\s+', ' ', str(text or '').strip().lower())


def _header_map(row):
    """{field: column index} when this row looks like a header (it has a name and a price column)."""
    cells = [_key(c) for c in row]
    found = {}
    for field, names in COLUMNS.items():
        for i, cell in enumerate(cells):
            if cell in names and i not in found.values():
                found[field] = i
                break
    return found if 'name' in found and 'price' in found else None


def _price(value):
    if value is None or value == '':
        return None
    text = re.sub(r'(?i)rs\.?|inr|₹|,|\s', '', str(value))
    try:
        price = Decimal(text)
    except InvalidOperation:
        return None
    return price.quantize(Decimal('0.01')) if price >= 0 else None


def _yes_no(value, default=None):
    text = _key(value)
    if text in YES:
        return True
    if text in NO:
        return False
    return default


def _int(value):
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def _sheets(content, filename):
    """[(sheet name, rows)] from the file."""
    name = filename.lower()
    if name.endswith('.csv'):
        text = content.decode('utf-8-sig', errors='replace')
        return [('Menu', list(csv.reader(io.StringIO(text))))]
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:  # a damaged or password-protected file
            raise MenuSheetError(f'Could not open {filename}: {error.__class__.__name__}.') from error
        return [(sheet.title, [list(r) for r in sheet.iter_rows(values_only=True)]) for sheet in book.worksheets
                if sheet.sheet_state == 'visible']
    raise MenuSheetError(f'{filename}: only .xlsx and .csv menus can be read automatically.')


def read_menu(content, filename):
    """
    {'items': [...], 'categories': [names in order], 'problems': [text], 'sheets_used': n}.
    Each item: name, category, price (Decimal) and any optional fields that were present.
    """
    items, problems, categories, used = [], [], [], 0
    skipped = 0

    for sheet_name, rows in _sheets(content, filename):
        header_at, columns = None, None
        for i, row in enumerate(rows[:15]):
            columns = _header_map(row)
            if columns:
                header_at = i
                break
        if columns is None:
            continue
        used += 1
        heading = None
        for n, row in enumerate(rows[header_at + 1:], start=header_at + 2):
            cell = lambda field: row[columns[field]] if field in columns and columns[field] < len(row) else None  # noqa: E731
            name = str(cell('name') or '').strip()
            if not name:
                continue
            price = _price(cell('price'))
            category = str(cell('category') or '').strip()
            if price is None:
                if 'category' not in columns and cell('price') in (None, ''):
                    heading = name.title() if name.isupper() else name  # a category heading row
                else:
                    skipped += 1
                    if len(problems) < MAX_PROBLEMS:
                        problems.append(f'{sheet_name} row {n}: "{name}" has no usable price')
                continue
            category = category or heading or (sheet_name if len(rows) and sheet_name.lower() not in ('sheet1', 'menu') else 'Menu')
            item = {'name': name[:200], 'category': category[:100], 'price': price}
            for field in ('description', 'kitchen_notes', 'short_name'):
                if cell(field):
                    item[field] = str(cell(field)).strip()[:1000]
            for field in ('vegetarian', 'vegan', 'gluten_free', 'spicy'):
                flag = _yes_no(cell(field))
                if flag is not None:
                    item[field] = flag
            if 'available' in columns:
                item['available'] = _yes_no(cell('available'), default=True)
            if cell('allergens') and _key(cell('allergens')) not in ('no', 'none', '-', 'nil'):
                item['allergens'] = [a.strip() for a in re.split(r'[;,]', str(cell('allergens'))) if a.strip()][:20]
            for field in ('calories', 'prep_minutes', 'sort_order'):
                number = _int(cell(field))
                if number is not None and number >= 0:
                    item[field] = number
            if category.lower() not in (c.lower() for c in categories):
                categories.append(category)
            items.append(item)
            if len(items) >= MAX_ITEMS:
                problems.append(f'Only the first {MAX_ITEMS} items were read.')
                break

    if not used:
        raise MenuSheetError(f'{filename}: no header row with an item name and a price column was found '
                             '(e.g. "Name" and "Price").')
    if skipped > MAX_PROBLEMS:
        problems.append(f'…and {skipped - MAX_PROBLEMS} more rows without a price')
    return {'items': items, 'categories': categories, 'problems': problems, 'sheets_used': used}
